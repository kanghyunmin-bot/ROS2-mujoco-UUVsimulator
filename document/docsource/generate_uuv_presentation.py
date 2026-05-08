from __future__ import annotations

import json
import math
import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
DOC_DIR = ROOT / "document"
FIG_DIR = DOC_DIR / "figures"
OUT_PATH = DOC_DIR / "uuv_mujoco_presentation_kr.pptx"
SUMMARY_PATH = DOC_DIR / "measurement_summary_latest.json"

FONT = os.environ.get("UUV_DOC_FONT", "Apple SD Gothic Neo")
FONT_FILE = os.environ.get("UUV_DOC_FONT_FILE", "")
if not FONT_FILE:
    for candidate in (
        "/System/Library/Fonts/AppleSDGothicNeo.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    ):
        if Path(candidate).exists():
            FONT_FILE = candidate
            break
FONT_FILE = FONT_FILE or None
COLOR_BG = RGBColor(255, 255, 255)
COLOR_TEXT = RGBColor(24, 31, 38)
COLOR_MUTED = RGBColor(89, 101, 114)
COLOR_ACCENT = RGBColor(15, 86, 122)
COLOR_ACCENT_2 = RGBColor(194, 110, 36)
COLOR_ACCENT_3 = RGBColor(55, 124, 84)
COLOR_PANEL = RGBColor(249, 251, 253)
COLOR_BORDER = RGBColor(216, 224, 232)


def load_summary() -> dict:
    return json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))["summary"]


def set_slide_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG


def add_title(slide, text: str, subtitle: str | None = None) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(12.1), Inches(0.72))
    tf = title_box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(23)
    run.font.bold = True
    run.font.color.rgb = COLOR_TEXT
    p.alignment = PP_ALIGN.LEFT
    try:
        tf.fit_text(font_family=FONT, max_size=23, bold=True, font_file=FONT_FILE)
    except Exception:
        pass
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.62), Inches(0.9), Inches(12.0), Inches(0.34))
        tf = sub_box.text_frame
        tf.clear()
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = subtitle
        run.font.name = FONT
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_MUTED
        try:
            tf.fit_text(font_family=FONT, max_size=10, font_file=FONT_FILE)
        except Exception:
            pass


def add_footer(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.65), Inches(7.04), Inches(12.0), Inches(0.18))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(8.5)
    run.font.color.rgb = COLOR_MUTED


def add_bullets(slide, bullets: list[str], left: float, top: float, width: float, height: float,
                font_size: float = 16, color: RGBColor = COLOR_TEXT, level0_space_after: float = 4) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(2)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.clear()
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {bullet}"
        p.level = 0
        p.space_after = Pt(level0_space_after)
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.1
        for run in p.runs:
            run.font.name = FONT
            run.font.size = Pt(font_size)
            run.font.color.rgb = color
    try:
        tf.fit_text(font_family=FONT, max_size=int(font_size), font_file=FONT_FILE)
    except Exception:
        pass


def add_panel(slide, left: float, top: float, width: float, height: float, title: str | None = None):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_PANEL
    line = shape.line
    line.color.rgb = COLOR_BORDER
    line.width = Pt(1.0)
    if title:
        title_box = slide.shapes.add_textbox(Inches(left + 0.18), Inches(top + 0.11), Inches(width - 0.36), Inches(0.28))
        p = title_box.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = title
        run.font.name = FONT
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_ACCENT
    return shape


def add_panel_text(slide, text: str, left: float, top: float, width: float, height: float,
                   font_size: float = 14, color: RGBColor = COLOR_TEXT, bold: bool = False,
                   align=PP_ALIGN.LEFT) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    tf.margin_top = Pt(3)
    tf.margin_bottom = Pt(2)
    tf.clear()
    for idx, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
        p.alignment = align
        p.space_after = Pt(2)
    try:
        tf.fit_text(font_family=FONT, max_size=int(font_size), bold=bold, font_file=FONT_FILE)
    except Exception:
        pass


def add_image(slide, path: Path, left: float, top: float, width: float | None = None, height: float | None = None):
    kwargs = {}
    if width is not None:
        kwargs["width"] = Inches(width)
    if height is not None:
        kwargs["height"] = Inches(height)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), **kwargs)


def metric_text(summary: dict) -> dict[str, str]:
    roll_deg = summary["stabilize_forward_step"]["roll_rad"]["mean"] * 180.0 / math.pi
    pitch_deg = summary["stabilize_forward_step"]["pitch_rad"]["mean"] * 180.0 / math.pi
    return {
        "manual_forward": f"MANUAL forward\nmean {summary['manual_forward_step']['surge_mps']['mean']:.3f} m/s\nmax {summary['manual_forward_step']['surge_mps']['max']:.3f} m/s",
        "stabilize_forward": f"STABILIZE forward\nmean {summary['stabilize_forward_step']['surge_mps']['mean']:.3f} m/s\nmax {summary['stabilize_forward_step']['surge_mps']['max']:.3f} m/s",
        "alt_hold_forward": f"ALT_HOLD forward\nmean {summary['alt_hold_forward_step']['surge_mps']['mean']:.3f} m/s\nmax {summary['alt_hold_forward_step']['surge_mps']['max']:.3f} m/s",
        "yaw": f"MANUAL yaw\nmean {summary['manual_yaw_step']['yaw_rate_rad_s']['mean']:.3f} rad/s\nfinal {summary['manual_yaw_step']['yaw_rate_rad_s']['final']:.3f} rad/s",
        "heave": f"ALT_HOLD heave\ndepth Δ {summary['alt_hold_heave_step']['depth_m']['delta']:.3f} m",
        "attitude": f"STABILIZE 자세\nroll {roll_deg:.2f} deg\npitch {pitch_deg:.2f} deg",
    }


def make_slide_1(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "MuJoCo 기반 UUV 시뮬레이터", "실시간 제어 스택 연동과 수중 물리 고도화")

    add_panel(slide, 0.72, 1.35, 6.3, 4.85, "문제 정의")
    add_bullets(
        slide,
        [
            "ArduSub/ROS2 제어 스택과 연결되는 실시간 UUV 시뮬레이터가 필요함",
            "기존 단순 모델은 부력, 선형 drag, 단순 thrust 위주라 실제 수중 응답과 차이가 큼",
            "핵심 목표는 엔진 비교가 아니라 수중 물리 항을 직접 설계하고 제어 응답까지 검증하는 환경을 만드는 것",
        ],
        left=0.95,
        top=1.82,
        width=5.8,
        height=3.95,
        font_size=17,
    )

    add_panel(slide, 7.25, 1.35, 5.35, 4.85, "핵심 키워드")
    add_panel_text(
        slide,
        "Real-time UUV simulation\n\nMuJoCo + custom hydrodynamics\n\nArduSub SITL / ROS2 integration\n\nJoystick / GUI / RC override",
        7.55,
        1.95,
        4.7,
        3.7,
        font_size=18,
        color=COLOR_TEXT,
    )
    add_footer(slide, "Slide 1 / 5")


def make_slide_2(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "원래 방식과 현재 커스텀 방식", "기성 fluid model / plugin 사용과 현재 runtime 구조의 차이")

    add_panel(slide, 0.72, 1.35, 5.55, 2.45, "원래는 어떻게 하는가")
    add_bullets(
        slide,
        [
            "MuJoCo라면 density, viscosity, fluidshape=ellipsoid, fluidcoef로 built-in fluid model을 사용하는 것이 정석",
            "Gazebo라면 보통 수중 plugin이나 hydrodynamics plugin 같은 기성 플러그인 구조를 사용",
        ],
        left=0.98,
        top=1.78,
        width=5.0,
        height=1.48,
        font_size=14.2,
    )

    add_panel(slide, 0.72, 4.05, 5.55, 2.5, "이번 프로젝트는 어떻게 커스텀했는가")
    add_bullets(
        slide,
        [
            "MuJoCo는 rigid-body solver로만 사용",
            "Python runtime이 수중 힘과 토크를 계산",
            "계산된 hydrodynamic wrench를 xfrc_applied로 주입",
            "최근에는 equivalent ellipsoid 기반 baseline coefficient도 생성",
        ],
        left=0.98,
        top=4.45,
        width=5.0,
        height=1.72,
        font_size=14.2,
    )

    add_panel(slide, 6.5, 1.35, 6.1, 5.2, "왜 이렇게 커스텀했는가")
    add_bullets(
        slide,
        [
            "제어 응답에 중요한 added mass, restoring torque, quadratic damping, thruster asymmetry를 직접 넣기 위해",
            "블랙박스 plugin이 아니라 항 하나하나의 물리적 의미와 응답 변화를 직접 해석하기 위해",
            "ArduSub/ROS2와 연결된 closed-loop testbed를 만들기 위해",
        ],
        left=6.78,
        top=1.76,
        width=5.55,
        height=1.9,
        font_size=14.4,
    )
    add_image(slide, FIG_DIR / "sim_real_hydrodynamic_coeffs.png", 6.78, 3.8, width=5.55)
    add_footer(slide, "Slide 2 / 5")


def make_slide_3(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "sim2real gap을 줄이기 위해 추가한 물리 항", "수중 물리와 추진기 모델을 실제 응답 쪽으로 보정")

    add_panel(slide, 0.72, 1.35, 5.15, 5.75, "추가 및 강화한 항목")
    add_bullets(
        slide,
        [
            "CoB restoring torque로 roll/pitch 복원 특성 반영",
            "Added mass, added-mass Coriolis, quadratic damping 추가",
            "Relative-current drag로 물에 대한 상대속도 기반 힘 계산",
            "T200 공개 성능표 기반 비선형 thrust curve 사용",
            "Reverse asymmetry, thruster별 gain calibration 반영",
            "Equivalent ellipsoid에서 baseline coefficient 생성",
        ],
        left=0.98,
        top=1.76,
        width=4.65,
        height=4.95,
        font_size=15.2,
    )

    add_panel(slide, 6.05, 1.35, 6.55, 2.7, "T200 thrust curve")
    add_image(slide, FIG_DIR / "thruster_curve_comparison.png", 6.25, 1.73, width=6.15)

    add_panel(slide, 6.05, 4.28, 6.55, 2.82, "Thruster calibration")
    add_image(slide, FIG_DIR / "thruster_gain_calibration.png", 6.25, 4.64, width=6.15)
    add_footer(slide, "Slide 3 / 5")


def make_slide_4(prs: Presentation, summary: dict):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "실제 ROS bag 기반 검증 결과", "mode 전환과 RC step 입력에 대한 실제 응답 측정")
    metrics = metric_text(summary)

    add_panel(slide, 0.72, 1.35, 8.15, 5.72, "Measured step response")
    add_image(slide, FIG_DIR / "actual_mavros_step_responses.png", 0.94, 1.72, width=7.7)

    add_panel(slide, 9.0, 1.35, 3.6, 3.55, "대표 수치")
    add_panel_text(
        slide,
        f"{metrics['manual_forward']}\n\n{metrics['stabilize_forward']}\n\n{metrics['alt_hold_forward']}",
        9.22,
        1.72,
        3.15,
        2.85,
        font_size=14.2,
        bold=True,
        color=COLOR_TEXT,
    )

    add_panel(slide, 9.0, 5.05, 3.6, 2.02, "Yaw / Heave / 해석")
    add_panel_text(
        slide,
        f"{metrics['yaw']}\n\n{metrics['heave']}\n\nSTABILIZE와 ALT_HOLD에서도 전진 응답이 유지됨",
        9.22,
        5.38,
        3.12,
        1.38,
        font_size=12.8,
        color=COLOR_TEXT,
    )
    add_footer(slide, "Slide 4 / 5")


def make_slide_5(prs: Presentation, summary: dict):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "이 접근의 이점과 다음 단계", "커스텀 물리 계층을 올렸을 때 얻는 연구적 의미")
    metrics = metric_text(summary)

    add_panel(slide, 0.72, 1.35, 5.15, 5.75, "이를 통해 얻을 수 있는 이점")
    add_bullets(
        slide,
        [
            "해석 가능성: 어떤 물리 항이 어떤 응답 변화를 만드는지 설명 가능",
            "확장성: 형상이 바뀌어도 ellipsoid baseline으로 새 tuning의 출발점 생성 가능",
            "검증 가능성: ROS mode 전환과 RC 입력을 실제 bag으로 기록해 closed-loop 검증 가능",
            "sim2real 개선성: thrust curve, damping, added mass, gain을 항목별로 식별 가능",
        ],
        left=0.98,
        top=1.78,
        width=4.6,
        height=4.95,
        font_size=15.2,
    )

    add_panel(slide, 6.0, 1.35, 6.6, 2.7, "Mode별 전진 응답")
    add_image(slide, FIG_DIR / "actual_mavros_mode_comparison.png", 6.22, 1.72, width=6.15)

    add_panel(slide, 6.0, 4.28, 2.95, 2.82, "XY trajectory")
    add_image(slide, FIG_DIR / "actual_mavros_xy_track.png", 6.18, 4.6, width=2.58)

    add_panel(slide, 9.2, 4.28, 3.4, 2.82, "다음 단계")
    add_panel_text(
        slide,
        "실기 velocity / yaw / depth 로그와\n현재 simulator 응답을 직접 비교해\ncoefficient identification을 진행하고,\ncurrent / turbulence / wave 모델을 확장할 계획입니다.",
        9.42,
        4.62,
        2.95,
        1.35,
        font_size=14.0,
        color=COLOR_TEXT,
    )
    add_panel_text(
        slide,
        f"{metrics['heave']}\n{metrics['attitude']}",
        9.42,
        6.0,
        2.95,
        0.62,
        font_size=11.0,
        color=COLOR_MUTED,
    )
    add_footer(slide, "Slide 5 / 5")


def build_presentation() -> None:
    summary = load_summary()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    make_slide_1(prs)
    make_slide_2(prs)
    make_slide_3(prs)
    make_slide_4(prs, summary)
    make_slide_5(prs, summary)

    prs.save(str(OUT_PATH))
    print(f"saved {OUT_PATH}")


if __name__ == "__main__":
    build_presentation()
